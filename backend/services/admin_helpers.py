import csv
import io
import logging
import os
import re
import secrets
import string
from collections import defaultdict

from fastapi import APIRouter, Depends, HTTPException, UploadFile
from models.sql_models import SqlBatch, SqlCourse, SqlUser
from services.academic_service import AcademicService
from sqlalchemy.exc import IntegrityError

from routes_fastapi.deps import get_current_user

logger = logging.getLogger("CampusGrade.AdminRoutes")
router = APIRouter(tags=["admin"])

def require_admin(current_user=Depends(get_current_user)):
    """Sprint 5, the original implementation plan: "any admin" gate (both `super_admin`
    and `program_admin`) for the ~12 routes in `routes_fastapi/admin/users.py`
    that use `admin["_id"]`/`admin.get(...)` dict access for audit logging
    and `created_by` â€” those call sites can't switch to the `Principal`
    dataclass `core.policy.require(...)` returns without touching every one
    of them. Kept dict-in/dict-out, but the allowed-role check now reads
    from `core.policy.PERMISSIONS[Permission.USER_MANAGE]` â€” the single
    source of truth â€” rather than a second, hand-maintained role set here."""
    from core.policy import PERMISSIONS, Permission
    from core.principal import Role

    try:
        role = Role(current_user.get("role"))
    except ValueError:
        role = None
    if role not in PERMISSIONS[Permission.USER_MANAGE]:
        raise HTTPException(status_code=403, detail="Admin access required")
    return current_user


def _lock_user_safety_context(session, user_id):
    """Serialize changes that could remove the final active super admin.

    Sprint 5, the original implementation plan: the continuity guard is specifically
    about `super_admin` (there must always be â‰¥1 active one) â€” a
    `program_admin` being deleted or demoted never trips this, since a
    program admin's absence doesn't leave the platform unmanageable the way
    losing the last super admin would."""
    active_admins = session.query(SqlUser).filter(
        SqlUser.role == "super_admin", SqlUser.is_active.is_(True),
    ).order_by(SqlUser.id).with_for_update().all()
    user = session.query(SqlUser).filter(SqlUser.id == user_id).with_for_update().first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    return user, active_admins


def _assert_admin_continuity(user, active_admins, actor_id, *, next_role=None, next_active=None, deleting=False):
    next_role = user.role if next_role is None else next_role
    next_active = user.is_active if next_active is None else next_active
    removes_active_admin = bool(
        user.role == "super_admin" and user.is_active
        and (deleting or next_role != "super_admin" or not next_active)
    )
    self_removal = str(user.id) == str(actor_id) and (
        deleting or (user.is_active and not next_active) or (user.role == "super_admin" and next_role != "super_admin")
    )
    if self_removal:
        raise HTTPException(status_code=409, detail="You cannot remove or deactivate your own administrator account")
    if removes_active_admin and len(active_admins) <= 1:
        raise HTTPException(status_code=409, detail="At least one active super administrator account is required")


def _commit_delete(session, entity_label):
    """Commit a delete while translating FK protection into a usable API response."""
    try:
        session.commit()
    except IntegrityError as exc:
        session.rollback()
        logger.info("Blocked %s delete because dependent records exist: %s", entity_label, exc)
        raise HTTPException(
            status_code=409,
            detail=f"Cannot delete this {entity_label} while it is still used by other records",
        ) from exc


def _normalize_user_email(raw_email):
    email = str(raw_email or "").strip().lower()
    if not email:
        raise HTTPException(status_code=400, detail="Email is required")
    if len(email) > 254:
        raise HTTPException(status_code=400, detail="Email cannot exceed 254 characters")
    return email


def _normalize_person_name(raw_name):
    name = " ".join(str(raw_name or "").strip().split())
    if len(name) < 2:
        raise HTTPException(status_code=400, detail="Name must contain at least 2 characters")
    if len(name) > 100:
        raise HTTPException(status_code=400, detail="Name cannot exceed 100 characters")
    return name


def _normalize_contact_number(raw_contact_number, *, required=False):
    raw_value = str(raw_contact_number or "").strip()
    if not raw_value:
        if required:
            raise HTTPException(status_code=400, detail="Contact number is required")
        return None
    digits = re.sub(r"\D", "", raw_value)
    if not 10 <= len(digits) <= 15:
        raise HTTPException(status_code=400, detail="Contact number must contain 10 to 15 digits")
    return f"+{digits}" if raw_value.startswith("+") else digits


def _normalize_student_roll_no(raw_roll_no):
    roll_no = re.sub(r"\s+", "", str(raw_roll_no or "").strip().upper())
    if not roll_no:
        return None
    if len(roll_no) > 50:
        raise HTTPException(status_code=400, detail="Roll number cannot exceed 50 characters")
    if not re.fullmatch(r"[A-Z0-9-]+", roll_no):
        raise HTTPException(status_code=400, detail="Roll number can only contain letters, numbers, and hyphens")
    return roll_no


def _normalize_teacher_id(raw_teacher_id):
    teacher_id = str(raw_teacher_id or "").strip().upper()
    if not teacher_id:
        return None

    teacher_id = re.sub(r"\s+", "", teacher_id)
    if len(teacher_id) > 50:
        raise HTTPException(status_code=400, detail="Teacher ID cannot exceed 50 characters")
    if not re.fullmatch(r"[A-Z0-9-]+", teacher_id):
        raise HTTPException(status_code=400, detail="Teacher ID can only contain letters, numbers, and hyphens")

    return teacher_id


def _generate_teacher_id_from_user_id(user_id):
    compact = str(user_id or "").replace("-", "").upper()
    if not compact:
        raise HTTPException(status_code=400, detail="Unable to generate teacher ID")
    return f"TCH-{compact[:8]}"


# `_generate_teacher_password`/`_generate_student_password` and their
# `_extract_first_name_for_password`/`_extract_last_two_digits` helpers were
# deleted here (Sprint 4, the original implementation plan): `firstname@last2digits`
# of a phone number or roll number is guessable by any classmate who knows
# that person's public info. Every call site (bulk teacher/student import,
# here and in `routes_fastapi/admin/users.py`) now uses this function.
def _generate_temporary_password():
    """Create a strong one-time password without persisting its plaintext."""
    return "".join((
        secrets.choice(string.ascii_uppercase),
        secrets.choice(string.ascii_lowercase),
        secrets.choice(string.digits),
        "!",
        secrets.token_urlsafe(12),
    ))


def _normalize_optional_fk(raw_value):
    value = str(raw_value or "").strip()
    return value or None


def _validate_student_batch_link(session, course_id=None, batch_id=None):
    if not batch_id:
        return None
    return AcademicService.require_batch_for_course(session, batch_id, course_id)


def _validate_student_section_link(session, section_id=None, course_id=None, batch_id=None):
    if not section_id:
        return None
    return AcademicService.require_section_for_batch(session, section_id, course_id, batch_id)

IMPORT_HEADER_ALIASES = {
    "name": {"name", "full_name", "teacher_name", "student_name"},
    "email": {"email", "email_id", "mail"},
    "teacher_id": {"teacher_id", "teacherid", "faculty_id", "employee_id"},
    "contact_number": {"contact_number", "contact", "phone", "phone_number", "mobile", "mobile_number"},
    "designation": {"designation", "title"},
    "roll_no": {"roll_no", "roll", "roll_number", "rollnumber", "admission_no"},
    "course": {"course", "program", "course_name", "program_name"},
    "batch": {"batch", "batch_name"},
    "section": {"section", "section_name"},
    "section_id": {"section_id"},
    "course_id": {"course_id", "program_id"},
    "batch_id": {"batch_id"},
    "year": {"year"},
}
MAX_IMPORT_FILE_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 5_000
MAX_IMPORT_COLUMNS = 50
IMPORT_READ_CHUNK_BYTES = 64 * 1024


def _coerce_import_cell_value(raw_value):
    if raw_value is None:
        return ""
    if isinstance(raw_value, float) and raw_value.is_integer():
        return str(int(raw_value))
    return str(raw_value).strip()


def _normalize_import_header(raw_header):
    header = _coerce_import_cell_value(raw_header).lower()
    header = re.sub(r"[^a-z0-9]+", "_", header)
    return header.strip("_")


def _canonicalize_import_header(raw_header):
    normalized_header = _normalize_import_header(raw_header)
    if not normalized_header:
        return ""

    for canonical_name, aliases in IMPORT_HEADER_ALIASES.items():
        if normalized_header in aliases:
            return canonical_name

    return normalized_header


def _parse_csv_import_bytes(file_bytes):
    try:
        text_content = file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text_content = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            raise HTTPException(status_code=400, detail="CSV must be UTF-8 encoded") from None

    reader = csv.DictReader(io.StringIO(text_content))
    if not reader.fieldnames:
        raise HTTPException(status_code=400, detail="CSV file must include a header row")

    if len(reader.fieldnames) > MAX_IMPORT_COLUMNS:
        raise HTTPException(status_code=413, detail=f"Import files may contain at most {MAX_IMPORT_COLUMNS} columns")

    rows = []
    for row_number, row in enumerate(reader, start=2):
        if not row:
            continue
        has_any_value = any(_coerce_import_cell_value(value) for value in row.values())
        if not has_any_value:
            continue
        rows.append((row_number, row))
        if len(rows) > MAX_IMPORT_ROWS:
            raise HTTPException(status_code=413, detail=f"Import files may contain at most {MAX_IMPORT_ROWS} data rows")

    return list(reader.fieldnames), rows


def _parse_excel_import_bytes(file_bytes):
    try:
        from openpyxl import load_workbook
    except Exception:
        raise HTTPException(status_code=500, detail="Excel import requires openpyxl installed on backend") from None

    if not file_bytes.startswith(b"PK"):
        raise HTTPException(status_code=400, detail="Excel file is not a valid .xlsx archive")

    try:
        workbook = load_workbook(filename=io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Excel file could not be read") from exc
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if not header_row:
            raise HTTPException(status_code=400, detail="Excel sheet is empty")

        headers = [_coerce_import_cell_value(value) for value in header_row]
        if len(headers) > MAX_IMPORT_COLUMNS:
            raise HTTPException(status_code=413, detail=f"Import files may contain at most {MAX_IMPORT_COLUMNS} columns")
        rows = []
        for row_number, value_row in enumerate(iterator, start=2):
            if not value_row:
                continue
            values = [_coerce_import_cell_value(cell_value) for cell_value in value_row]
            if not any(values):
                continue
            raw_row = {
                headers[index] if index < len(headers) else f"column_{index + 1}": values[index]
                for index in range(len(values))
            }
            rows.append((row_number, raw_row))
            if len(rows) > MAX_IMPORT_ROWS:
                raise HTTPException(status_code=413, detail=f"Import files may contain at most {MAX_IMPORT_ROWS} data rows")

        return headers, rows
    finally:
        workbook.close()


def _parse_import_file(file_name, file_bytes):
    extension = os.path.splitext(str(file_name or "").strip().lower())[1]
    if extension == ".csv":
        return _parse_csv_import_bytes(file_bytes)
    if extension == ".xlsx":
        return _parse_excel_import_bytes(file_bytes)
    raise HTTPException(status_code=400, detail="Unsupported file type. Use .csv or .xlsx")


async def _read_import_file(file: UploadFile):
    extension = os.path.splitext(str(file.filename or "").strip().lower())[1]
    if extension not in {".csv", ".xlsx"}:
        raise HTTPException(status_code=400, detail="Unsupported file type. Use .csv or .xlsx")

    chunks = []
    total_bytes = 0
    while True:
        chunk = await file.read(IMPORT_READ_CHUNK_BYTES)
        if not chunk:
            break
        total_bytes += len(chunk)
        if total_bytes > MAX_IMPORT_FILE_BYTES:
            raise HTTPException(status_code=413, detail="Import file exceeds the 5MB limit")
        chunks.append(chunk)

    if not chunks:
        raise HTTPException(status_code=400, detail="Uploaded file is empty")
    return b"".join(chunks)


def _normalize_import_row(raw_row):
    normalized_row = {}
    for raw_key, raw_value in (raw_row or {}).items():
        canonical_key = _canonicalize_import_header(raw_key)
        if not canonical_key:
            continue

        value = _coerce_import_cell_value(raw_value)
        if canonical_key not in normalized_row or not normalized_row.get(canonical_key):
            normalized_row[canonical_key] = value

    return normalized_row


def _validate_required_import_columns(raw_headers, required_columns):
    normalized_headers = {
        _canonicalize_import_header(header)
        for header in (raw_headers or [])
        if _canonicalize_import_header(header)
    }
    missing_columns = [column for column in required_columns if column not in normalized_headers]
    if missing_columns:
        raise HTTPException(
            status_code=400,
            detail=f"Missing required columns: {', '.join(missing_columns)}",
        )


def _safe_positive_int(raw_value, field_name):
    if raw_value is None:
        return None

    value_text = str(raw_value).strip()
    if not value_text:
        return None

    try:
        parsed_value = int(value_text)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail=f"{field_name} must be a number") from None

    if parsed_value < 1 or parsed_value > 12:
        raise HTTPException(status_code=400, detail=f"{field_name} must be between 1 and 12")

    return parsed_value


def _build_import_response(message, total_rows, added_rows, duplicate_rows, invalid_rows, dry_run=False):
    return {
        "message": message,
        "dry_run": bool(dry_run),
        "summary": {
            "total_rows": total_rows,
            "added": len(added_rows),
            "duplicates": len(duplicate_rows),
            "invalid": len(invalid_rows),
        },
        "added_rows": added_rows,
        "duplicate_rows": duplicate_rows,
        "invalid_rows": invalid_rows,
    }

@router.get("/users")
def _normalize_program_name(raw_name):
    name = " ".join(str(raw_name or "").strip().split())
    if not name:
        raise HTTPException(status_code=400, detail="Program name is required")
    if len(name) > 100:
        raise HTTPException(status_code=400, detail="Program name cannot exceed 100 characters")
    return name


def _normalize_lookup_text(raw_value):
    return " ".join(str(raw_value or "").strip().lower().split())


def _compute_teacher_allocation_counts(session):
    """Count distinct teachers allocated to each program/batch/section/subject.

    Teacher teaching allocations are stored as a JSON list on each teacher
    row (``SqlUser.assignments``), not a join table, so this has to scan
    every active teacher's allocations rather than run a SQL count. Fine at
    the scale of one school; revisit if this ever becomes a bottleneck.
    """
    program_teachers = defaultdict(set)
    batch_teachers = defaultdict(set)
    section_teachers = defaultdict(set)
    subject_teachers = defaultdict(set)
    teachers = session.query(SqlUser.id, SqlUser.assignments).filter(SqlUser.role == "teacher", SqlUser.is_active.is_(True)).all()
    for teacher_id, allocations in teachers:
        for allocation in allocations or []:
            if allocation.get("program_id"):
                program_teachers[allocation["program_id"]].add(teacher_id)
            if allocation.get("batch_id"):
                batch_teachers[allocation["batch_id"]].add(teacher_id)
            if allocation.get("section_id"):
                section_teachers[allocation["section_id"]].add(teacher_id)
            for subject_id in allocation.get("subject_ids") or []:
                subject_teachers[subject_id].add(teacher_id)
    return {
        "program": program_teachers,
        "batch": batch_teachers,
        "section": section_teachers,
        "subject": subject_teachers,
    }


def _normalize_batch_lookup_text(raw_value):
    normalized = _normalize_lookup_text(raw_value)
    if not normalized:
        return ""
    return re.sub(r"\s*-\s*", "-", normalized)


def _build_course_aliases(course_name, course_code):
    aliases = set()
    name = str(course_name or "").strip()
    code = str(course_code or "").strip()

    normalized_name = _normalize_lookup_text(name)
    normalized_code = _normalize_lookup_text(code)

    if normalized_name:
        aliases.add(normalized_name)
    if normalized_code:
        aliases.add(normalized_code)
    if normalized_name and normalized_code:
        aliases.add(_normalize_lookup_text(f"{name} ({code})"))

    return aliases


def _build_batch_aliases(batch):
    aliases = set()
    aliases.add(_normalize_batch_lookup_text(_build_batch_label(batch.start_year, batch.end_year, batch.name)))
    aliases.add(_normalize_batch_lookup_text(f"{batch.start_year}-{batch.end_year}"))
    aliases.add(_normalize_batch_lookup_text(f"{batch.start_year} - {batch.end_year}"))
    return {alias for alias in aliases if alias}


# --- Course CRUD ---
def _normalize_course_year(raw_year):
    try:
        year = int(raw_year)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="Program year must be a valid number") from None

    if year < 1 or year > 12:
        raise HTTPException(status_code=400, detail="Program year must be between 1 and 12")
    return year


def _normalize_course_code(raw_code):
    code = str(raw_code or "").strip().upper()
    if not code:
        raise HTTPException(status_code=400, detail="Program code is required")
    if len(code) > 50:
        raise HTTPException(status_code=400, detail="Program code cannot exceed 50 characters")
    return code


# --- Batch CRUD ---
def _normalize_batch_year(raw_year, field_name):
    try:
        year = int(raw_year)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail=f"{field_name} must be a valid number") from None

    if year < 1900 or year > 2500:
        raise HTTPException(status_code=400, detail=f"{field_name} must be between 1900 and 2500")
    return year


def _normalize_batch_name(raw_name):
    name = " ".join(str(raw_name or "").strip().split())
    if not name:
        return None
    if len(name) > 100:
        raise HTTPException(status_code=400, detail="Batch name cannot exceed 100 characters")
    return name


def _build_batch_label(start_year, end_year, name=None):
    return str(name or "").strip() or f"{start_year} - {end_year}"


def _normalize_section_name(raw_name):
    name = " ".join(str(raw_name or "").strip().split())
    if not name:
        raise HTTPException(status_code=400, detail="Section name is required")
    if len(name) > 100:
        raise HTTPException(status_code=400, detail="Section name cannot exceed 100 characters")
    return name


# --- Section CRUD ---
def _validate_section_links(session, course_id, batch_id):
    course = session.query(SqlCourse).filter(SqlCourse.id == course_id).first()
    if not course:
        raise HTTPException(status_code=404, detail="Program not found")

    batch = session.query(SqlBatch).filter(SqlBatch.id == batch_id).first()
    if not batch:
        raise HTTPException(status_code=404, detail="Batch not found")
    if batch.course_id != course_id:
        raise HTTPException(status_code=400, detail="Batch does not belong to selected program")


# --- Subject CRUD ---
def _normalize_subject_name(raw_name):
    name = " ".join(str(raw_name or "").strip().split())
    if not name:
        raise HTTPException(status_code=400, detail="Course name is required")
    if len(name) > 100:
        raise HTTPException(status_code=400, detail="Course name cannot exceed 100 characters")
    return name


def _normalize_subject_code(raw_code):
    code = str(raw_code or "").strip().upper()
    if not code:
        raise HTTPException(status_code=400, detail="Course code is required")
    if len(code) > 50:
        raise HTTPException(status_code=400, detail="Course code cannot exceed 50 characters")
    return code


def _normalize_subject_credits(raw_credits):
    try:
        credits = int(raw_credits)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="Credit must be a valid number") from None

    if credits < 1 or credits > 20:
        raise HTTPException(status_code=400, detail="Credit must be between 1 and 20")
    return credits


def _normalize_subject_year(raw_year, default_year=1):
    if raw_year is None:
        return default_year

    try:
        year = int(raw_year)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="Course year must be a valid number") from None

    if year < 1 or year > 12:
        raise HTTPException(status_code=400, detail="Course year must be between 1 and 12")
    return year


def _normalize_subject_fk(raw_value, field_label, required=True):
    value = str(raw_value or "").strip()
    if not value:
        if required:
            raise HTTPException(status_code=400, detail=f"{field_label} is required")
        return None
    return value
